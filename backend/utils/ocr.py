import fitz  # PyMuPDF
import docx
import openpyxl
from PIL import Image
import io

def extract_text_from_pdf(file_bytes: bytes) -> list[dict]:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages = []
    for page_num, page in enumerate(doc):
        text = page.get_text().strip()
        if text:
            pages.append({"page": page_num + 1, "text": text, "source": "native"})
    doc.close()
    return pages

def extract_text_from_docx(file_bytes: bytes) -> list[dict]:
    doc = docx.Document(io.BytesIO(file_bytes))
    chunks = []
    current = []
    page_num = 1
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            current.append(text)
        if len(current) >= 20:
            chunks.append({"page": page_num, "text": "\n".join(current), "source": "native"})
            page_num += 1
            current = []
    if current:
        chunks.append({"page": page_num, "text": "\n".join(current), "source": "native"})
    return chunks

def extract_text_from_excel(file_bytes: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    pages = []
    for sheet_num, sheet in enumerate(wb.worksheets):
        rows = []
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                rows.append(row_text)
        if rows:
            pages.append({
                "page": sheet_num + 1,
                "text": f"Sheet: {sheet.title}\n" + "\n".join(rows),
                "source": "native"
            })
    return pages

def extract_text_from_txt(file_bytes: bytes) -> list[dict]:
    text = file_bytes.decode("utf-8", errors="ignore").strip()
    # Split into chunks of ~100 lines
    lines = text.split("\n")
    chunks = []
    size = 100
    for i in range(0, len(lines), size):
        chunk = "\n".join(lines[i:i+size]).strip()
        if chunk:
            chunks.append({"page": i // size + 1, "text": chunk, "source": "native"})
    return chunks